"""CSV and XLSX bytes from export rows. Text cells are already safe (`rows.safe_text`)."""

import csv
import io
from collections.abc import Sequence
from datetime import date
from decimal import Decimal

from openpyxl import Workbook

from app.export.rows import COLUMNS, Cell

CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _csv_cell(value: Cell) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")  # Indian order
    return value


def to_csv(rows: Sequence[dict[str, Cell]]) -> bytes:
    """UTF-8 with a BOM so Excel shows the rupee sign and names correctly."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([heading for _, heading, _ in COLUMNS])
    for row in rows:
        writer.writerow([_csv_cell(row.get(key)) for key, _, _ in COLUMNS])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_xlsx(rows: Sequence[dict[str, Cell]]) -> bytes:
    """Money as real numbers (2 decimals) and dates as real dates, so sums and filters work."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Approved"
    sheet.append([heading for _, heading, _ in COLUMNS])
    for row in rows:
        sheet.append([row.get(key) for key, _, _ in COLUMNS])
    for column, (_, heading, kind) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = max(
            14, len(heading) + 2
        )
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            if kind == "money":
                cell.number_format = "#,##0.00"
            elif kind == "date":
                cell.number_format = "dd-mm-yyyy"
            elif isinstance(cell.value, str):
                cell.data_type = "s"  # never a formula, whatever the text starts with
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
