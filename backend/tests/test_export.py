import io
from datetime import date
from decimal import Decimal as D

from openpyxl import load_workbook

from app.clients import CLIENTS_BY_ID
from app.export.files import to_csv, to_xlsx
from app.export.rows import COLUMNS, ExportFacts, build_row, safe_text
from app.normalisation.dates import parse_date
from app.schemas.documents import Invoice, LineItem, Receipt, TaxLine
from app.schemas.issues import IssueCode, make_issue

CLIENT = CLIENTS_BY_ID["C001"]
EVIL = '=HYPERLINK("http://evil.example","click")'


def invoice(**changes) -> Invoice:
    base = {
        "vendor_name": "Sai Stationery Mart",
        "vendor_gstin": "23QSSPM6142B1ZF",
        "invoice_number": "INV/26-27/001",
        "invoice_date": parse_date("14-08-2026"),
        "place_of_supply_code": "23",
        "line_items": [LineItem(description="Paper", taxable_amount=D("1000"))],
        "taxable_value": D("1000"),
        "tax_lines": [
            TaxLine(head="CGST", rate=D(9), amount=D("90")),
            TaxLine(head="SGST", rate=D(9), amount=D("90")),
        ],
        "round_off": D("-0.20"),
        "total": D("1179.80"),
    }
    return Invoice(**{**base, **changes})


def facts(**changes) -> ExportFacts:
    base = {
        "document_id": "doc-1",
        "client": CLIENT,
        "original_filename": "bill.pdf",
        "ledger": {"selected_code": "6150", "suggested_code": "6150", "overridden": False},
        "issues": [],
        "approved_at": "2026-11-01T10:00:00+00:00",
    }
    return ExportFacts(**{**base, **changes})


def test_safe_text_defuses_every_formula_start_and_leaves_other_text_alone():
    for start in ("=", "+", "-", "@", "\t", "\r"):
        assert safe_text(f"{start}1+1").startswith("'" + start)
    assert safe_text("Sai Stationery") == "Sai Stationery"
    assert safe_text("") == ""


def test_an_invoice_row_has_sums_ledger_name_and_only_warnings():
    issues = [
        make_issue(IssueCode.MISSING_HSN_SAC, "no HSN"),
        make_issue(IssueCode.TOTAL_MISMATCH, "an error that was fixed or never exported"),
    ]
    row = build_row(invoice(), facts(issues=issues))
    assert row["cgst"] == D("90") and row["sgst"] == D("90") and row["igst"] is None
    assert row["total"] == D("1179.80") and row["round_off"] == D("-0.20")
    assert (row["ledger_code"], row["ledger_name"]) == ("6150", "Printing and stationery")
    assert row["warnings"] == "missing_hsn_sac"
    assert row["date"] == date(2026, 8, 14) and row["client_name"] == CLIENT.name
    assert row["ledger_overridden"] == "no"
    assert set(row) == {key for key, _, _ in COLUMNS}


def test_a_receipt_row_has_no_place_of_supply_or_taxable_value():
    receipt = Receipt(
        merchant_name="Shree Fuels",
        receipt_date=parse_date("03-09-2026"),
        total=D("1250.50"),
    )
    row = build_row(receipt, facts())
    assert row["document_type"] == "receipt" and row["vendor_name"] == "Shree Fuels"
    assert row["place_of_supply"] is None and row["taxable_value"] is None


def test_text_from_the_bill_is_made_safe_in_the_row_but_numbers_are_not_touched():
    row = build_row(invoice(vendor_name=EVIL, invoice_number="-5"), facts(original_filename="@x"))
    assert row["vendor_name"] == "'" + EVIL
    assert row["number"] == "'-5" and row["original_filename"] == "'@x"
    assert row["round_off"] == D("-0.20")  # a negative amount stays a number


def test_csv_has_a_bom_headings_indian_dates_two_decimals_and_safe_text():
    data = to_csv([build_row(invoice(vendor_name=EVIL), facts())])
    assert data.startswith(b"\xef\xbb\xbf")
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0].split(",")[0] == "Document ID" and len(lines) == 2
    assert "14-08-2026" in lines[1] and "1179.80" in lines[1] and "-0.20" in lines[1]
    assert "'=HYPERLINK" in lines[1] and ",=HYPERLINK" not in lines[1]


def test_csv_with_no_rows_is_just_the_headings():
    assert len(to_csv([]).decode("utf-8-sig").splitlines()) == 1


def test_xlsx_has_real_numbers_real_dates_and_text_that_is_never_a_formula():
    book = load_workbook(io.BytesIO(to_xlsx([build_row(invoice(vendor_name=EVIL), facts())])))
    sheet = book["Approved"]
    headings = [c.value for c in sheet[1]]
    row = {h: c for h, c in zip(headings, sheet[2], strict=True)}
    assert row["Total"].value == 1179.80 and row["Total"].number_format == "#,##0.00"
    assert row["Round off"].value == -0.20
    assert row["Date"].value.date() == date(2026, 8, 14)
    vendor = row["Vendor / merchant"]
    assert vendor.value == "'" + EVIL and vendor.data_type == "s"
    assert sheet.max_row == 2 and sheet.freeze_panes == "A2"
