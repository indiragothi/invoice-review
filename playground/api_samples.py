"""Experiment: drive the HTTP API with FastAPI's TestClient. Offline, no API calls.

A temporary data folder and the cached model answers stand in for Gemini. Checks every route that
exists so far and its status codes (docs/api-and-pipeline.md). Run from backend/:

    uv run --locked --no-sync python ../playground/api_samples.py
"""

import csv
import io
import json
import shutil
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BACKEND))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pymupdf  # noqa: E402
from email_samples import cached_draft  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pipeline_samples import (  # noqa: E402
    cached_extractor,
    cached_ledger_suggester,
    cached_second_reader,
)

from app.config import APP_CONFIG  # noqa: E402
from app.main import create_app  # noqa: E402
from app.providers import CallInfo, ExtractionResult, ProviderError  # noqa: E402

WORK = Path(__file__).resolve().parent / "output" / "api_check"
results: list[bool] = []


def check(label: str, ok: bool, detail: str = "") -> None:
    results.append(ok)
    print(f"{'OK ' if ok else 'BAD'} {label}{(': ' + detail) if detail else ''}")


def main() -> None:
    manifest = json.loads((APP_CONFIG.samples_dir / "manifest.json").read_text(encoding="utf-8"))
    documents = manifest["documents"]
    by_stem = {Path(d["file"]).stem: d for d in documents}
    mode = {"stem": "", "fail": None}  # which cached answer to use; how to misbehave

    def extractor(reading):
        if mode["fail"] == "busy":
            raise ProviderError("The AI provider is busy. Retry later.", busy=True)
        result = cached_extractor(mode["stem"])(reading)
        if mode["fail"] == "other":
            other = result.extraction.model_copy(update={"document_type": "other"})
            return ExtractionResult(extraction=other, call=CallInfo(**result.call.__dict__))
        return result

    def ledger(document, client_name):
        return cached_ledger_suggester(mode["stem"])(document, client_name)

    shutil.rmtree(WORK, ignore_errors=True)

    def drafter(_document, _client, _issues):
        if mode["fail"] == "email":
            raise ProviderError("The AI provider is busy. Retry later.", busy=True)
        return cached_draft(mode["stem"], "email_v1")

    def second_reader(images, document_type):
        result = cached_second_reader(mode["stem"])(images, document_type)
        if mode["fail"] == "tamper":  # a reading B that disagrees about the total
            changed = result.reading.model_copy(update={"total": "99,999.00"})
            return type(result)(reading=changed, call=result.call)
        return result

    client = TestClient(create_app(WORK, extractor, ledger, drafter, second_reader))

    def upload(path_or_bytes, name, client_id):
        data = path_or_bytes.read_bytes() if isinstance(path_or_bytes, Path) else path_or_bytes
        return client.post(
            "/api/documents", files={"file": (name, data)}, data={"client_id": client_id}
        )

    r = client.get("/health")
    check("GET /health", r.status_code == 200 and r.json()["status"] == "ok")
    r = client.get("/api/clients")
    check("GET /api/clients", r.status_code == 200 and len(r.json()) == 3, r.json()[0]["name"])
    r = client.get("/api/accounting/ledgers")
    check("GET /api/accounting/ledgers", r.status_code == 200 and len(r.json()) == 13)
    r = client.options(
        "/api/documents",
        headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST"},
    )
    check(
        "CORS allows the Vite origin",
        r.headers.get("access-control-allow-origin") == "http://localhost:5173",
    )

    # upload all 12
    ids, status_ok = {}, 0
    for entry in documents:
        path = APP_CONFIG.samples_dir / entry["file"]
        mode.update(stem=path.stem, fail=None)
        r = upload(path, path.name, entry["client"])
        ok = r.status_code == 201 and r.json()["status"] == entry["expected_status"]
        status_ok += ok
        if r.status_code == 201:
            ids[path.stem] = r.json()["id"]
    check("POST /api/documents x12 -> 201 and manifest status", status_ok == 12, f"{status_ok}/12")

    body = client.get(f"/api/documents/{ids['01-intra-classic']}").json()
    check(
        "GET /api/documents/{id}",
        body["status"] == "ready"
        and body["ledger"]["suggested_code"] == "6150"
        and "reading" not in body,
        f"ledger {body['ledger']['suggested_code']}, keys {len(body)}",
    )
    ready = client.get("/api/documents", params={"status": "ready"}).json()
    review = client.get("/api/documents", params={"status": "needs_review"}).json()
    c001 = client.get("/api/documents", params={"client_id": "C001"}).json()
    check(
        "GET /api/documents filters",
        (len(ready), len(review), len(c001)) == (8, 4, 10),
        f"ready {len(ready)}, needs_review {len(review)}, C001 {len(c001)}",
    )
    check(
        "list is newest first",
        [d["id"] for d in client.get("/api/documents").json()][0] == ids["12-restaurant-receipt"],
    )
    check(
        "bad status filter -> 422",
        client.get("/api/documents", params={"status": "nope"}).status_code == 422,
    )

    sample = APP_CONFIG.samples_dir / by_stem["01-intra-classic"]["file"]
    r = client.get(f"/api/documents/{ids['01-intra-classic']}/file")
    check(
        "GET .../file returns the original bytes",
        r.content == sample.read_bytes()
        and r.headers["content-type"] == "application/pdf"
        and r.headers["x-content-type-options"] == "nosniff",
    )
    r = client.get(f"/api/documents/{ids['01-intra-classic']}/pages/1")
    check(
        "GET .../pages/1 is a PNG", r.status_code == 200 and r.content[:8] == b"\x89PNG\r\n\x1a\n"
    )
    check(
        "page 2 of a 1-page file -> 404",
        client.get(f"/api/documents/{ids['01-intra-classic']}/pages/2").status_code == 404,
    )
    check(
        "page 2 of the 2-page invoice -> 200",
        client.get(f"/api/documents/{ids['09-two-page']}/pages/2").status_code == 200,
    )
    check("unknown document -> 404", client.get("/api/documents/nope").status_code == 404)

    # rejected uploads
    mode["stem"] = "01-intra-classic"
    before = len(client.get("/api/documents").json())
    pdf = pymupdf.open()
    for _ in range(APP_CONFIG.max_pages + 1):
        pdf.new_page()
    six_pages = pdf.tobytes()
    cases = {
        "renamed .exe -> 415": (upload(b"MZ" + b"\0" * 80, "bill.pdf", "C001"), 415),
        "over 5 MB -> 413": (
            upload(b"%PDF-" + b"0" * APP_CONFIG.max_upload_bytes, "big.pdf", "C001"),
            413,
        ),
        "6 pages -> 413": (upload(six_pages, "long.pdf", "C001"), 413),
        "unknown client -> 404": (upload(sample, "x.pdf", "C999"), 404),
    }
    for label, (response, want) in cases.items():
        check(label, response.status_code == want, response.json().get("detail", ""))
    check("rejections store nothing", len(client.get("/api/documents").json()) == before)
    check(
        "a path in the file name is not used",
        upload(sample, "../../evil.pdf", "C001").status_code == 201
        and not any("evil" in p.name for p in (WORK / "uploads").iterdir()),
    )

    # failures that are saved
    other = APP_CONFIG.samples_dir / by_stem["04-inter-compact"]["file"]
    client.delete(f"/api/documents/{ids['04-inter-compact']}")
    mode.update(stem="04-inter-compact", fail="busy")
    r = upload(other, other.name, by_stem["04-inter-compact"]["client"])
    failed_id = r.json()["document"]["id"]
    check(
        "provider failure -> 502, record saved as failed",
        r.status_code == 502
        and client.get(f"/api/documents/{failed_id}").json()["status"] == "failed",
        r.json()["detail"],
    )
    mode["fail"] = None
    check(
        "retry a non-failed document -> 409",
        client.post(f"/api/documents/{ids['01-intra-classic']}/retry").status_code == 409,
    )
    r = client.post(f"/api/documents/{failed_id}/retry")
    check(
        "retry the failed one -> 200 ready", r.status_code == 200 and r.json()["status"] == "ready"
    )
    mode["fail"] = "other"
    r = upload(sample, "notbill.pdf", "C002")
    check(
        "not an invoice or receipt -> 422, saved as failed",
        r.status_code == 422 and r.json()["document"]["status"] == "failed",
        r.json()["detail"],
    )
    failed_doc = r.json()["document"]["id"]
    mode["fail"] = None

    # ledger selection, decision, bulk approve: the approval invariant
    def put_ledger(doc_id, code):
        return client.put(f"/api/documents/{doc_id}/ledger", json={"ledger_code": code})

    def decide(doc_id, decision):
        return client.post(f"/api/documents/{doc_id}/decision", json={"decision": decision})

    a, b, c = ids["01-intra-classic"], ids["09-two-page"], ids["03-inter-classic"]
    bad = ids["05-invalid-vendor-gstin"]  # needs_review
    r = decide(a, "approved")
    check(
        "approve a ready document with no ledger selected -> 409",
        r.status_code == 409 and "ledger" in r.json()["detail"],
        r.json()["detail"],
    )
    check("unknown ledger code -> 404", put_ledger(a, "9999").status_code == 404)
    r = put_ledger(a, "6150")
    check(
        "confirm the suggestion -> overridden false",
        r.status_code == 200
        and r.json()["ledger"]["selected_code"] == "6150"
        and not r.json()["ledger"]["overridden"],
    )
    r = put_ledger(a, "6190")
    check(
        "pick another ledger -> overridden true, history kept",
        r.json()["ledger"]["overridden"] and r.json()["history"][-1]["event"] == "ledger_selected",
    )
    put_ledger(a, "6150")
    r = decide(a, "approved")
    check(
        "approve a ready document with a ledger -> 200 approved",
        r.status_code == 200 and r.json()["status"] == "approved",
    )
    r = decide(a, "approved")
    check("approve again -> 409 locked", r.status_code == 409 and "locked" in r.json()["detail"])
    check("reject an approved record -> 409", decide(a, "rejected").status_code == 409)
    check("change the ledger of a decided record -> 409", put_ledger(a, "6190").status_code == 409)
    check(
        "approved stays approved", client.get(f"/api/documents/{a}").json()["status"] == "approved"
    )

    put_ledger(bad, "6140")
    r = decide(bad, "approved")
    check(
        "approve needs_review even with a ledger -> 409",
        r.status_code == 409
        and "total" not in r.json()["detail"]
        and "needs_review" in r.json()["detail"],
        r.json()["detail"],
    )
    check(
        "the same call fails again, nothing changed",
        client.get(f"/api/documents/{bad}").json()["status"] == "needs_review",
    )
    check("approve a failed document -> 409", decide(failed_doc, "approved").status_code == 409)
    check(
        "select a ledger on a failed document -> 409",
        put_ledger(failed_doc, "6190").status_code == 409,
    )
    check(
        "bad decision value -> 422",
        client.post(f"/api/documents/{b}/decision", json={"decision": "maybe"}).status_code == 422,
    )
    check("decision on an unknown document -> 404", decide("nope", "approved").status_code == 404)

    put_ledger(b, "5100")
    r = client.post(
        "/api/documents/bulk-approve", json={"ids": [b, c, bad, a, failed_doc, "nope", b]}
    )
    body = r.json()
    reasons = {s["id"]: s["reason"] for s in body["skipped"]}
    check(
        "bulk approve: only the ready one with a ledger is approved",
        r.status_code == 200
        and [d["id"] for d in body["approved"]] == [b]
        and body["approved"][0]["status"] == "approved",
        f"approved {len(body['approved'])}, skipped {len(body['skipped'])}",
    )
    check(
        "bulk approve: the others are skipped with a reason",
        set(reasons) == {c, bad, a, failed_doc, "nope"}
        and "ledger" in reasons[c]
        and "needs_review" in reasons[bad]
        and "locked" in reasons[a]
        and reasons["nope"] == "Document not found.",
    )
    check(
        "bulk approve: the invalid ids list is bounded",
        client.post("/api/documents/bulk-approve", json={"ids": ["x"] * 201}).status_code == 422,
    )

    # reject, and a rejected record no longer counts as a duplicate
    r = decide(c, "rejected")
    check(
        "reject a ready document -> 200 rejected",
        r.status_code == 200 and r.json()["status"] == "rejected",
    )
    check("reject a needs_review document -> 200", decide(bad, "rejected").status_code == 200)
    path_03 = APP_CONFIG.samples_dir / by_stem["03-inter-classic"]["file"]
    mode["stem"] = "03-inter-classic"
    r = upload(path_03, "03-again.pdf", by_stem["03-inter-classic"]["client"])
    codes = [i["code"] for i in r.json()["issues"]]
    check(
        "re-upload of a rejected bill: duplicate_file warns, duplicate_invoice does not",
        "duplicate_file" in codes
        and "duplicate_invoice" not in codes
        and r.json()["status"] == "ready",
    )
    check(
        "the export rule starts here: approved documents",
        [
            d["status"] for d in client.get("/api/documents", params={"status": "approved"}).json()
        ].count("approved")
        == 2,
    )

    # field edits: review_data, source human, re-validation, status moves both ways
    def edit(doc_id, fields):
        return client.put(f"/api/documents/{doc_id}", json={"fields": fields})

    def codes_of(response):
        return [i["code"] for i in response.json()["issues"]]

    mode.update(stem="05-invalid-vendor-gstin", fail=None)
    r = upload(
        APP_CONFIG.samples_dir / by_stem["05-invalid-vendor-gstin"]["file"], "05.pdf", "C001"
    )
    fix_id = r.json()["id"]
    fields_before = r.json()["review_data"]["fields"]
    check(
        "review_data lists every field with value, source, confidence and evidence box",
        r.json()["status"] == "needs_review"
        and "vendor_gstin" in fields_before
        and fields_before["vendor_gstin"]["evidence"]["box"] is not None
        and fields_before["vendor_gstin"]["source"] == "pdf_text",
        f"{len(fields_before)} fields",
    )
    wrong = fields_before["vendor_gstin"]["value"]
    gstin_issue = next(i for i in r.json()["issues"] if i["code"] == "invalid_vendor_gstin")
    good = gstin_issue["suggestion"]  # the one-character fix the app suggests
    assert good and good != wrong, gstin_issue
    r = edit(fix_id, {"vendor_gstin": good})
    entry = r.json()["review_data"]["fields"]["vendor_gstin"]
    check(
        "fix the GSTIN -> human, confidence 1.0, original kept, status ready",
        r.status_code == 200
        and r.json()["status"] == "ready"
        and entry["source"] == "human"
        and entry["confidence"] == 1.0
        and entry["original"] == wrong
        and "invalid_vendor_gstin" not in codes_of(r),
        f"{wrong} -> {good}",
    )
    last = r.json()["history"][-1]
    check(
        "the edit is in the history with from / to",
        last["event"] == "edited"
        and last["changes"]["vendor_gstin"] == {"from": wrong, "to": good},
    )
    r = edit(fix_id, {"vendor_gstin": good})
    check(
        "an unchanged value changes nothing",
        r.status_code == 200
        and len(r.json()["history"])
        == len(client.get(f"/api/documents/{fix_id}").json()["history"]),
    )
    put_ledger(fix_id, "6140")
    r = decide(fix_id, "approved")
    check(
        "edit -> select ledger -> approve works end to end",
        r.status_code == 200 and r.json()["status"] == "approved",
    )
    check(
        "edit an approved record -> 409", edit(fix_id, {"vendor_gstin": wrong}).status_code == 409
    )

    broken = ids["06-total-mismatch"]
    printed = client.get(f"/api/documents/{broken}").json()["review_data"]["fields"]["total"][
        "value"
    ]
    r = edit(broken, {"total": "75,138.00"})
    check(
        "fix the total -> blocking error gone, status ready, words warning left",
        r.status_code == 200
        and r.json()["status"] == "ready"
        and "total_mismatch" not in codes_of(r)
        and "amount_in_words_mismatch" in codes_of(r),
        f"{printed} -> 75,138.00; {codes_of(r)}",
    )

    ok_doc = ids["02-intra-compact"]
    r = edit(ok_doc, {"total": "999.00"})
    check(
        "a wrong total on a ready document -> needs_review with total_mismatch",
        r.json()["status"] == "needs_review" and "total_mismatch" in codes_of(r),
    )
    original_total = r.json()["review_data"]["fields"]["total"]["original"]
    r = edit(ok_doc, {"total": original_total})
    check(
        "put the original value back -> ready again",
        r.json()["status"] == "ready" and "total_mismatch" not in codes_of(r),
    )
    r = edit(ok_doc, {"invoice_number": ""})
    check(
        "clearing the invoice number -> missing_invoice_number, needs_review",
        r.json()["status"] == "needs_review"
        and "missing_invoice_number" in codes_of(r)
        and r.json()["review_data"]["fields"]["invoice_number"]["grounding"] == "missing",
    )
    first = client.get(f"/api/documents/{a}").json()["review_data"]["fields"]
    r = edit(
        ok_doc,
        {
            "invoice_number": first["invoice_number"]["value"],
            "vendor_gstin": first["vendor_gstin"]["value"],
            "invoice_date": first["invoice_date"]["value"],
        },
    )
    check(
        "edit into another invoice's key -> duplicate_invoice is found",
        "duplicate_invoice" in codes_of(r) and r.json()["status"] == "needs_review",
        str(codes_of(r)),
    )

    n_issues = len(client.get(f"/api/documents/{ok_doc}").json()["issues"])
    r = edit(ok_doc, {"no_such_field": "x", "total": "1.00"})
    check(
        "unknown field -> 422 and nothing applied",
        r.status_code == 422
        and len(client.get(f"/api/documents/{ok_doc}").json()["issues"]) == n_issues,
        r.json()["detail"],
    )
    check(
        "line item that does not exist -> 422",
        edit(ok_doc, {"line_items[40].hsn_sac": "1"}).status_code == 422,
    )
    check(
        "value over 500 characters -> 422",
        edit(ok_doc, {"vendor_name": "x" * 501}).status_code == 422,
    )
    check("empty edit -> 422", edit(ok_doc, {}).status_code == 422)
    check("edit a failed document -> 409", edit(failed_doc, {"total": "1"}).status_code == 409)
    check("edit a rejected document -> 409", edit(c, {"total": "1"}).status_code == 409)
    check("edit an unknown document -> 404", edit("nope", {"total": "1"}).status_code == 404)

    # correction email: eligibility is fixed code; the draft is text only
    mode.update(stem="08-buyer-not-client", fail=None)
    r = upload(APP_CONFIG.samples_dir / by_stem["08-buyer-not-client"]["file"], "08.pdf", "C001")
    mail_id = r.json()["id"]
    flagged = [i for i in r.json()["issues"] if i["supplier_fixable"]]
    check(
        "issues say whether the supplier can fix them",
        r.json()["supplier_action_required"] is True
        and [i["code"] for i in flagged] == ["buyer_not_client"],
        str([(i["code"], i["supplier_fixable"]) for i in r.json()["issues"]]),
    )
    r = client.post(f"/api/documents/{mail_id}/correction-email")
    check(
        "draft the correction email -> 200 with subject, body and the eligible codes",
        r.status_code == 200
        and r.json()["issue_codes"] == ["buyer_not_client"]
        and "RP/RENT/26-27/07" in r.json()["subject"]
        and "23QXGPS4821K1ZU" in r.json()["body"],
        r.json()["subject"],
    )
    check(
        "the draft is recorded in the history, nothing is sent or stored",
        client.get(f"/api/documents/{mail_id}").json()["history"][-1]["event"]
        == "correction_email_drafted",
    )
    mode["fail"] = "email"
    check(
        "provider failure while drafting -> 502",
        client.post(f"/api/documents/{mail_id}/correction-email").status_code == 502,
    )
    mode["fail"] = None
    r = edit(mail_id, {"buyer_gstin": "23QXGPS4821K1ZU"})
    check(
        "fix the buyer GSTIN -> nothing left for the supplier, no email",
        r.json()["supplier_action_required"] is False
        and client.post(f"/api/documents/{mail_id}/correction-email").status_code == 409,
        r.json()["status"],
    )
    check(
        "a document with no supplier-fixable issue -> 409",
        client.post(f"/api/documents/{a}/correction-email").status_code == 409,
    )
    check(
        "a failed document -> 409",
        client.post(f"/api/documents/{failed_doc}/correction-email").status_code == 409,
    )
    check(
        "an unknown document -> 404",
        client.post("/api/documents/nope/correction-email").status_code == 404,
    )
    listing = client.get("/api/documents").json()
    check(
        "every listed document carries supplier_action_required",
        all(isinstance(d["supplier_action_required"], bool) for d in listing),
        f"{sum(d['supplier_action_required'] for d in listing)} of {len(listing)} need action",
    )

    # export: approved records only, safe for spreadsheets
    EVIL = '=HYPERLINK("http://evil.example","click")'
    r = edit(failed_id, {"vendor_name": EVIL})
    put_ledger(failed_id, "6190")
    check(
        "a bill whose vendor name is a formula can still be approved",
        r.status_code == 200 and decide(failed_id, "approved").status_code == 200,
    )

    approved = client.get("/api/documents", params={"status": "approved"}).json()
    approved_ids = {d["id"] for d in approved}
    r = client.get("/api/export")
    table = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    exported_ids = {row["Document ID"] for row in table}
    check(
        "GET /api/export (csv): only approved documents, as an attachment",
        r.status_code == 200
        and r.headers["content-type"].startswith("text/csv")
        and "attachment" in r.headers["content-disposition"]
        and exported_ids == approved_ids
        and len(table) == len(approved) >= 3,
        f"{len(table)} rows",
    )
    others = {d["id"] for d in client.get("/api/documents").json()} - approved_ids
    check(
        "no ready, needs_review, rejected or failed document is exported",
        not (exported_ids & others),
    )
    by_id = {row["Document ID"]: row for row in table}
    check(
        "the export shows the reviewer's corrected GSTIN and the chosen ledger",
        by_id[fix_id]["Vendor GSTIN"] == good
        and by_id[fix_id]["Ledger code"] == "6140"
        and by_id[fix_id]["Ledger"] == "Repairs and maintenance"
        and by_id[fix_id]["Ledger changed by reviewer"] == "no",
        f"{wrong} -> {by_id[fix_id]['Vendor GSTIN']}",
    )
    row_a = by_id[a]
    check(
        "amounts have two decimals, dates are DD-MM-YYYY, approval time is there",
        row_a["Total"].replace(".", "").isdigit()
        and row_a["Total"].split(".")[-1].__len__() == 2
        and len(row_a["Date"]) == 10
        and row_a["Date"][2] == "-"
        and row_a["Approved at"] != "",
        f"{row_a['Total']} on {row_a['Date']}",
    )
    check(
        "formula injection: the vendor name is exported as text",
        by_id[failed_id]["Vendor / merchant"] == "'" + EVIL,
        by_id[failed_id]["Vendor / merchant"][:30],
    )

    r = client.get("/api/export", params={"format": "xlsx"})
    book = load_workbook(io.BytesIO(r.content))
    sheet = book["Approved"]
    heads = [c.value for c in sheet[1]]
    cells = {row[heads.index("Document ID")].value: row for row in sheet.iter_rows(min_row=2)}
    vendor_cell = cells[failed_id][heads.index("Vendor / merchant")]
    total_cell = cells[a][heads.index("Total")]
    check(
        "GET /api/export?format=xlsx: real numbers, text that is never a formula",
        r.status_code == 200
        and set(cells) == approved_ids
        and isinstance(total_cell.value, int | float)
        and vendor_cell.value == "'" + EVIL
        and vendor_cell.data_type == "s",
        f"{len(cells)} rows, total {total_cell.value}",
    )
    c1 = client.get("/api/export", params={"client_id": "C001"}).content.decode("utf-8-sig")
    c2 = client.get("/api/export", params={"client_id": "C002"}).content.decode("utf-8-sig")
    approved_c001 = [d for d in approved if d["client_id"] == "C001"]
    check(
        "client filter; a client with nothing approved gets just the headings",
        len(c1.splitlines()) == 1 + len(approved_c001) and len(c2.splitlines()) == 1,
        f"C001 {len(c1.splitlines()) - 1} rows, C002 {len(c2.splitlines()) - 1}",
    )
    check(
        "unknown client -> 404",
        client.get("/api/export", params={"client_id": "C999"}).status_code == 404,
    )
    check(
        "unknown format -> 422",
        client.get("/api/export", params={"format": "pdf"}).status_code == 422,
    )
    check(
        "exporting changes nothing",
        {d["id"] for d in client.get("/api/documents", params={"status": "approved"}).json()}
        == approved_ids,
    )

    # second reading: a reading B that disagrees about the total (SECOND_READING=always)
    base_settings = client.app.state.settings
    fuel = APP_CONFIG.samples_dir / by_stem["11-fuel-receipt"]["file"]
    client.delete(f"/api/documents/{ids['11-fuel-receipt']}")  # so the upload is new again
    client.app.state.settings = base_settings.model_copy(update={"second_reading": "always"})
    mode.update(stem="11-fuel-receipt", fail="tamper")
    r = upload(fuel, "fuel.jpg", by_stem["11-fuel-receipt"]["client"])
    body = r.json()
    total = body["review_data"]["fields"]["total"]
    check(
        "a second reading that disagrees about the total -> needs_review, one warning",
        r.status_code == 201
        and body["status"] == "needs_review"
        and body["document_review"]["ran"] is True
        and body["document_review"]["agreement"]["total"] == "differ"
        and [i["code"] for i in body["issues"] if i["field"] == "total"] == ["readings_disagree"]
        and not any(i["severity"] == "error" for i in body["issues"]),
        str([i["code"] for i in body["issues"]]),
    )
    check(
        "the disputed value is capped at 0.5 and marked differ; the run record shows 3 calls",
        total["confidence"] == 0.5
        and total["agreement"] == "differ"
        and body["pipeline_run"]["calls"] == 3
        and body["pipeline_run"]["prompt_versions"]["review"] == "review_v1",
        f"confidence {total['confidence']}",
    )
    check(
        "readings_disagree is not something the supplier can fix",
        body["supplier_action_required"] is False
        and all(
            not i["supplier_fixable"] for i in body["issues"] if i["code"] == "readings_disagree"
        ),
    )
    fuel_id = body["id"]
    r = edit(fuel_id, {"total": total["original"]})
    fixed = r.json()
    check(
        "the reviewer confirms the value -> human, warning gone, ready",
        r.status_code == 200
        and fixed["status"] == "ready"
        and fixed["review_data"]["fields"]["total"]["agreement"] is None
        and fixed["review_data"]["fields"]["total"]["source"] == "human"
        and "readings_disagree"
        not in [i["code"] for i in fixed["issues"] if i["field"] == "total"],
    )
    client.delete(f"/api/documents/{fuel_id}")
    client.app.state.settings = base_settings.model_copy(update={"second_reading": "never"})
    mode["fail"] = None
    r = upload(fuel, "fuel.jpg", by_stem["11-fuel-receipt"]["client"])
    check(
        "SECOND_READING=never -> not run, 2 calls, status ready",
        r.json()["document_review"]["ran"] is False
        and r.json()["pipeline_run"]["calls"] == 2
        and r.json()["status"] == "ready"
        and r.json()["pipeline_run"]["prompt_versions"]["review"] is None,
    )
    client.delete(f"/api/documents/{r.json()['id']}")
    client.app.state.settings = base_settings

    # delete
    victim = ids["12-restaurant-receipt"]
    stored = {p.name for p in (WORK / "uploads").iterdir()}
    r = client.delete(f"/api/documents/{victim}")
    gone = client.get(f"/api/documents/{victim}").status_code == 404
    check(
        "DELETE -> 204, then 404, file removed",
        r.status_code == 204
        and gone
        and len({p.name for p in (WORK / "uploads").iterdir()}) == len(stored) - 1,
    )
    check("DELETE unknown -> 404", client.delete("/api/documents/nope").status_code == 404)

    print(f"\nAPI {sum(results)}/{len(results)} checks passed")


if __name__ == "__main__":
    main()
